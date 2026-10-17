"""Lectura local de Excel, con autorización previa y sin imprimir valores."""

from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, time
import hashlib
import warnings

import pandas as pd
from openpyxl import load_workbook

from sigard_ml.official.access import AccessDenied, AuthorizedUse, SourceAuthorization
from sigard_ml.official.settings import COLUMN_ALIASES, HEADER_SCAN_ROWS, MIN_HEADER_MATCHES, identify_column, safe_header


class ExcelReadError(ValueError):
    """Mensaje genérico: la excepción original puede contener datos sensibles."""


def file_digest(source: SourceAuthorization) -> str:
    source.require_source_use(AuthorizedUse.INSPECTION)
    with source.path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


@contextmanager
def open_authorized_workbook(source: SourceAuthorization, use: AuthorizedUse):
    source.require_source_use(use)
    if source.path.suffix.lower() != ".xlsx":
        raise ExcelReadError("Formato no habilitado; se requiere XLSX")
    workbook = None
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(source.path, read_only=True, data_only=False, keep_links=False)
            yield workbook
    except (AccessDenied, ExcelReadError):
        raise
    except Exception:
        raise ExcelReadError("No se pudo leer la fuente Excel; revisar localmente sin volcar celdas") from None
    finally:
        if workbook is not None:
            workbook.close()


def find_header(worksheet, configured_row=None) -> tuple[int, tuple]:
    if configured_row is not None:
        if type(configured_row) is not int or configured_row < 1:
            raise ExcelReadError("Fila de encabezado inválida")
        values = next(worksheet.iter_rows(min_row=configured_row, max_row=configured_row, values_only=True), ())
        if not values or all(is_empty(value) for value in values):
            raise ExcelReadError("Fila de encabezado vacía")
        return configured_row, tuple(values)
    for number, values in enumerate(worksheet.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True), 1):
        matches = {identify_column(value) for value in values} - {None}
        if len(matches) >= MIN_HEADER_MATCHES:
            return number, tuple(values)
    raise ExcelReadError("Encabezado no reconocido; configurar su fila explícitamente")


def read_sheet(source: SourceAuthorization, sheet: str, *, use=AuthorizedUse.PROCESSING,
               header_row=None, column_map=None) -> tuple[pd.DataFrame, dict]:
    source.require_use(sheet=sheet, subset_ref=source.subset_ref, use=use)
    with open_authorized_workbook(source, use) as workbook:
        if sheet not in workbook.sheetnames:
            raise ExcelReadError("Hoja solicitada inexistente")
        worksheet = workbook[sheet]
        row_number, headers = find_header(worksheet, header_row)
        column_map = column_map or {}
        labels, fields = [], {}
        for index, value in enumerate(headers, 1):
            canonical = identify_column(value)
            labels.append(f"column_{index:03d}")
            if canonical and canonical not in fields:
                fields[canonical] = labels[-1]
        for field, index in column_map.items():
            if field not in COLUMN_ALIASES:
                raise ExcelReadError("Campo configurado no admitido")
            if type(index) is not int or not 1 <= index <= len(labels):
                raise ExcelReadError("Índice de columna inválido")
            recognized = identify_column(headers[index - 1])
            if recognized is not None and recognized != field:
                raise ExcelReadError("El campo configurado contradice el encabezado reconocido")
            fields[field] = labels[index - 1]
        rows, excel_rows = [], []
        for excel_row, values in enumerate(worksheet.iter_rows(min_row=row_number + 1, values_only=True), row_number + 1):
            values = tuple(values[:len(labels)])
            if all(is_empty(value) for value in values):
                continue
            rows.append(values)
            excel_rows.append(excel_row)
        frame = pd.DataFrame(rows, columns=labels, dtype=object)
        frame.index = pd.Index(excel_rows, name="excel_row")
        return frame, {"header_row": row_number, "fields": fields, "columns": len(labels),
                       "safe_headers": {label: safe_header(value) for label, value in zip(labels, headers)}}


def value_type(value) -> str:
    if is_empty(value):
        return "empty"
    if isinstance(value, (date, datetime)):
        return "date"
    if isinstance(value, time):
        return "time_without_date"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def inspect_source(source: SourceAuthorization, sheet_options=None) -> dict:
    # No descubrir nombres de hojas de un libro con permisos parciales.
    source.require_source_use(AuthorizedUse.INSPECTION)
    if source.all_sheets_approved:
        with open_authorized_workbook(source, AuthorizedUse.INSPECTION) as workbook:
            sheets = tuple(workbook.sheetnames)
    else:
        sheets = source.sheets
    results = []
    for index, sheet in enumerate(sheets, 1):
        options = (sheet_options or {}).get(sheet, (sheet_options or {}).get(f"sheet_{index:02d}", {}))
        try:
            frame, schema = read_sheet(source, sheet, use=AuthorizedUse.INSPECTION,
                                       header_row=options.get("header_row"), column_map=options.get("column_map"))
        except ExcelReadError:
            results.append({"sheet_ref": f"sheet_{index:02d}", "status": "requires_header_review"})
            continue
        by_column = {column: field for field, column in schema["fields"].items()}
        columns = []
        for column in frame:
            types = Counter(value_type(value) for value in frame[column])
            columns.append({"column_ref": column, "field": by_column.get(column, "unmapped"),
                            "types": sorted(types), "empty_rows": types.get("empty", 0)})
        results.append({"sheet_ref": f"sheet_{index:02d}", "status": "inspected",
                        "rows": len(frame), **schema, "column_structure": columns})
    return {"source_id": source.source_id, "sheets": results}
