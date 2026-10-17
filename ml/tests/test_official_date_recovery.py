"""Fase 4.1: libros, fechas, identidades y registros completamente ficticios."""

from datetime import date
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from sigard_ml.official.date_recovery_pipeline import run_date_recovery
from sigard_ml.official.observed_comparison import compare_scenarios
from sigard_ml.official.observed_dates import (
    DATE_HIERARCHY, assign_record_weeks, identify_reference_header, parse_reference_date,
)
from sigard_ml.official.observed_pipeline import observed_output, run_observed_pipeline, validate_options
from sigard_ml.official.observed_reading import read_event_references
from sigard_ml.official.observed_settings import OBSERVED_DEFAULTS
from sigard_ml.official.pipeline import validate_delivery
from sigard_ml.official.reports import write_report
from sigard_ml.official.spatial_pipeline import digest_path, protected_snapshot, require_preserved, run_spatial_diagnostics
from sigard_ml.official.spatial_settings import SPATIAL_DEFAULTS

from test_official_observed import fixture_frame, options, panel_for, prepared, radios
from test_official_spatial import fictional_delivery


def temporal_frame(values):
    row = {}
    for field, value in zip(DATE_HIERARCHY, values, strict=True):
        row[field] = value
        row[field + "_status"] = "valid" if value is not None else "missing"
    return pd.DataFrame([row], dtype=object)


def assign(values, hierarchy=DATE_HIERARCHY, **changes):
    limits = dict(period_start=date(2020, 1, 1), period_end=date(2024, 12, 31),
                  as_of=date(2025, 1, 31), version="fictional-hierarchy-v1", **changes)
    return assign_record_weeks(temporal_frame(values), field="symptom_onset", hierarchy=hierarchy, **limits).iloc[0]


@pytest.mark.parametrize("values, kind, approximate", [
    ([date(2024, 1, 5), date(2024, 1, 6), date(2024, 1, 7), date(2024, 1, 8)], "inicio_sintomas", False),
    ([None, date(2024, 1, 6), date(2024, 1, 7), date(2024, 1, 8)], "consulta_medica", True),
    ([None, None, date(2024, 1, 7), date(2024, 1, 8)], "toma_muestra", True),
    ([None, None, None, date(2024, 1, 8)], "notificacion", True),
])
def test_first_valid_date_and_provenance(values, kind, approximate):
    result = assign(values)
    assert result.tipo_fecha_referencia == kind
    assert bool(result.fecha_aproximada) == approximate
    assert bool(result.semana_potencialmente_afectada_por_retraso) == approximate
    assert result.fecha_referencia == result.reference_date == next(v for v in values if v is not None)
    assert result.reference_date_usable


@pytest.mark.parametrize("invalid, status", [
    ("31/02/2024", "invalid_format"), ("2024-13-01", "invalid_format"),
    ("01/02/24", "ambiguous_date"), (45100, "invalid_format"),
    ("=TODAY()", "invalid_format"), (True, "invalid_format"),
    (date(1800, 1, 1), "out_of_range"), ("sin datos", "invalid_format"),
])
def test_invalid_values_are_unavailable_and_next_valid_field_is_selected(invalid, status):
    result = assign([invalid, None, date(2024, 1, 7), None])
    assert result.tipo_fecha_referencia == "toma_muestra"
    assert f"symptom_onset:{status}" in result.date_selection_rejections


def test_conflicting_explicit_formats_are_ambiguous_without_guessing():
    assert parse_reference_date("02/03/2024", ("%d/%m/%Y", "%m/%d/%Y")) == (None, "ambiguous_date")
    assert parse_reference_date("02/03/2024") == (date(2024, 3, 2), "valid")


@pytest.mark.parametrize("values", [
    [None] * 4, ["invalid"] * 4,
    [date(2026, 1, 1)] * 4, [date(2019, 1, 1)] * 4,
])
def test_no_valid_date_has_no_week_and_no_artificial_reference(values):
    result = assign(values)
    assert not result.reference_date_usable and not result.fecha_aproximada
    assert result.fecha_referencia is None and result.week_start is None
    assert result.motivo_seleccion_fecha == "ninguna_fecha_valida"


def test_outside_period_onset_is_skipped_before_sample():
    result = assign([date(2019, 12, 31), None, date(2020, 1, 3), None])
    assert result.tipo_fecha_referencia == "toma_muestra"
    assert "symptom_onset:outside_analysis_period" in result.date_selection_rejections


@pytest.mark.parametrize("day, year, week, start", [
    (date(2021, 1, 2), 2020, 53, date(2020, 12, 27)),
    (date(2021, 1, 3), 2021, 1, date(2021, 1, 3)),
    (date(2023, 12, 31), 2024, 1, date(2023, 12, 31)),
    (date(2024, 12, 29), 2025, 1, date(2024, 12, 29)),
])
def test_fallback_uses_existing_epidemiological_calendar(day, year, week, start):
    result = assign([None, None, day, None])
    assert (result.epidemiological_year, result.epidemiological_week, result.week_start) == (year, week, start)


@pytest.mark.parametrize("header", ["Fecha de revisión", "Fecha última actualización", "Fecha apertura", "Consulta geocodificador",
                                  "Fecha consulta", "Fecha muestra"])
def test_administrative_dates_never_identified_as_medical_consultation(header):
    assert identify_reference_header(header) is None


@pytest.mark.parametrize("policy", [["updated_on"], ["sample_on", "symptom_onset"], [], ["symptom_onset", "notified_on"]])
def test_unidentified_or_reordered_fields_are_rejected(policy):
    with pytest.raises(ValueError, match="jerarquía"):
        assign([None] * 4, hierarchy=policy)


def test_new_policy_requires_versioned_output_and_rule(tmp_path):
    with pytest.raises(ValueError, match="versionada"):
        observed_output(options(date_hierarchy=list(DATE_HIERARCHY)), tmp_path, True)
    with pytest.raises(ValueError, match="versión"):
        validate_options(options(date_hierarchy=list(DATE_HIERARCHY)))


def test_other_pending_reasons_are_preserved_and_counts_reconcile(radios):
    frame = fixture_frame(6)
    frame["fecha_inicio_sintomas_original"] = ["sin datos"] * 6
    frame["tipo_valor_fecha_inicio_sintomas_original"] = "text"
    frame[["symptom_onset", "symptom_onset_status"]] = [None, "missing"]
    for field in DATE_HIERARCHY[1:]:
        frame[field], frame[field + "_status"] = None, "missing"
    frame["sample_on"], frame["sample_on_status"] = date(2024, 1, 10), "valid"
    frame.loc[1, "event_ref"] = frame.loc[2, "event_ref"]
    frame.loc[3, "classification_label"] = "unknown fictional label"
    frame.loc[4, "radio_candidate"] = None
    frame.loc[5, "classification_label"] = "Descartado"
    before = frame.copy(deep=True)
    a, _ = prepared(frame, radios)
    b_settings = options(output_dir="data/interim/official/observed/fixture-expanded", date_hierarchy=list(DATE_HIERARCHY), date_rule_version="fixture-expanded-v1")
    b, conservative = prepared(frame, radios, settings=b_settings)
    assert b.processing_outcome.tolist() == ["included", "pending", "pending", "pending", "pending", "excluded"]
    comparison = compare_scenarios(a, b)
    assert comparison["recovered_by_date_type"]["toma_muestra"] == 1
    assert comparison["reconciliation_verified"]
    assert comparison["expanded"] == {"included": 1, "pending": 4, "excluded": 1, "represented_weeks": 1}
    panel = panel_for(b, radios, settings=b_settings)
    assert int(panel.registros_elegibles_provisionales.sum()) == 1
    assert not panel.registros_elegibles_provisionales.eq(0).any()
    assert b.fecha_inicio_sintomas_original.tolist() == ["sin datos"] * 6
    repeated, _ = prepared(frame, radios, settings=b_settings)
    pd.testing.assert_frame_equal(b, repeated)
    pd.testing.assert_frame_equal(frame, before)


@pytest.fixture
def recovery_delivery(fictional_delivery, monkeypatch):
    root, config, registry, source = fictional_delivery
    book = load_workbook(source)
    sheet = book.active
    for column, header in enumerate(["Fecha de consulta médica", "Fecha de toma de muestra", "Fecha de notificación",
                                    "Fecha de revisión administrativa", "Fecha apertura"], 8):
        sheet.cell(1, column, header)
    for row in range(2, 5):
        sheet.cell(row, 3).value = "  sin datos  "
        sheet.cell(row, 9, date(2024, 1, 7))
        sheet.cell(row, 11, date(2024, 1, 6))
        sheet.cell(row, 12, date(2024, 1, 6))
    book.save(source)
    book.close()
    write_report(validate_delivery(config, registry), root / "data/interim/official/quality_report.json")
    run_spatial_diagnostics(config, registry, dict(SPATIAL_DEFAULTS), root)
    monkeypatch.setattr("sigard_ml.official.observed_pipeline.require_branch", lambda root: None)
    monkeypatch.setattr("sigard_ml.official.date_recovery_pipeline.require_branch", lambda root: None)
    run_observed_pipeline(config, registry, dict(OBSERVED_DEFAULTS), root, exploratory=True)
    return root, config, registry, source


def test_excel_original_value_and_identified_date_columns(recovery_delivery):
    root, config, registry, source = recovery_delivery
    before = digest_path(source)
    records, schema = read_event_references(config, registry, root)
    assert records.fecha_inicio_sintomas_original.tolist() == ["  sin datos  "] * 3
    assert records.tipo_valor_fecha_inicio_sintomas_original.eq("text").all()
    assert records.consultation_on.isna().all() and records.notified_on.isna().all()
    assert records.sample_on.tolist() == [date(2024, 1, 7)] * 3
    assert set(schema[0]["identified_date_columns"]) == set(DATE_HIERARCHY)
    assert digest_path(source) == before


def test_complete_private_comparison_preserves_original_and_reproduces_all_files(recovery_delivery):
    root, config, registry, _ = recovery_delivery
    output = root / "data/interim/official/observed/date_recovery_041"
    before = protected_snapshot(root, output)
    report = run_date_recovery(config, registry, root, baseline_dir="data/interim/official/observed/exploratory")
    assert report["baseline_reproduced"] and report["baseline_preserved"]
    assert report["strict"]["included"] == 0 and report["expanded"]["included"] == 1
    assert report["radios"] == 263 and report["recovered_total"] == 1
    hashes = {str(path.relative_to(output)): digest_path(path) for path in output.rglob("*") if path.is_file()}
    run_date_recovery(config, registry, root, baseline_dir="data/interim/official/observed/exploratory")
    assert hashes == {str(path.relative_to(output)): digest_path(path) for path in output.rglob("*") if path.is_file()}
    require_preserved(before, protected_snapshot(root, output))
    metadata = json.loads((output / "expanded/exploratory/metadata.json").read_text(encoding="utf-8"))
    assert not metadata["alternative_dates_confirm_symptom_onset"] and not metadata["training_ready"]


def test_recovery_cannot_overwrite_original_or_change_eligibility(recovery_delivery):
    root, config, registry, _ = recovery_delivery
    baseline = "data/interim/official/observed/exploratory"
    for overrides in ({"output_dir": baseline}, {"output_dir": "data/processed/official"},
                      {"exploratory_include": ["probable"]}):
        with pytest.raises(ValueError):
            run_date_recovery(config, registry, root, baseline_dir=baseline, overrides=overrides)
    occupied = root / "data/interim/official/observed/occupied"
    occupied.mkdir()
    existing = occupied / "panel.parquet"
    existing.write_bytes(b"fictional-existing-panel")
    with pytest.raises(ValueError, match="ajenos"):
        run_date_recovery(config, registry, root, baseline_dir=baseline, overrides={"output_dir": str(occupied)})
    assert existing.read_bytes() == b"fictional-existing-panel"


def test_denied_authorization_stops_before_private_baseline_or_snapshot(recovery_delivery, monkeypatch):
    from dataclasses import replace
    from sigard_ml.official.access import AccessDenied, SourceRegistry
    root, config, registry, _ = recovery_delivery
    denied = SourceRegistry((replace(registry.sources[0], enabled=False),))
    monkeypatch.setattr("sigard_ml.official.date_recovery_pipeline.protected_snapshot",
                        lambda *args: pytest.fail("No debe leer archivos sin autorización"))
    with pytest.raises(AccessDenied):
        run_date_recovery(config, denied, root, baseline_dir="data/interim/official/observed/exploratory")
