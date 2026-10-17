"""Cotejos y XLSX ficticios: ningún dato real entra a las pruebas."""

from copy import deepcopy
from datetime import datetime
import json
from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from sigard_ml.official.selection_reconciliation import (
    VERSION, compare, date_value, identifier, profile, project_record, selection_state,
)
from sigard_ml.official.selection_reconciliation_pipeline import (
    cross_source_overlap, read_local_role, resolve_columns, run_reconciliation, validate_manifest,
)
from sigard_ml.official.spatial_pipeline import digest_path, protected_snapshot, require_preserved

KEY = b"fictional_test_only_key_32_bytes_!!"
COLUMNS = {"event_id": 1, "classification_columns": {"manual": 2}, "province": 3,
    "department": 4, "symptom_onset": 5, "address": 6}
RULE = {"classification_field": "manual", "accepted_classification_labels": ["Confirmado"],
    "province_values": ["La Rioja"], "department_values": ["Capital"],
    "reference_date_field": "symptom_onset", "require_valid_reference_date": None,
    "require_address_present": None}


def record(event, classification="Confirmado", province="La Rioja", department="Capital", day="2024-01-10", address="PRIVATE_ADDRESS", key=KEY):
    return project_record((event, classification, province, department, day, address), COLUMNS, key, "fictional")


def test_no_patient_values_or_identifiers_survive_projection():
    projected = record("PRIVATE_EVENT_ID")
    serialized = repr(projected)
    assert "PRIVATE_EVENT_ID" not in serialized and "PRIVATE_ADDRESS" not in serialized
    assert isinstance(projected["key"], bytes) and len(projected["key"]) == 32


def test_identifier_handling_preserves_leading_zeros_and_rejects_precision_loss():
    assert identifier("00123") == "00123"
    assert identifier(123) == identifier(123.0) == identifier("123")
    assert identifier(float(2**54)) is None
    assert identifier(True) is None and identifier("=A1") is None


def test_literal_rule_is_not_clinical_confirmation_and_does_not_require_dates_or_addresses():
    selected = record("a", day=None, address=None)
    before = deepcopy(selected)
    assert selection_state(selected, RULE)[0] is True
    state, failures, _ = selection_state(selected, {**RULE, "require_valid_reference_date": True, "require_address_present": True})
    assert state is False and len(failures) == 2
    assert selected == before
    assert selection_state(record("b", classification="Sospechoso"), RULE)[0] is False
    assert selection_state(record("c", classification=None), RULE)[0] is None


def test_formula_and_unmapped_address_are_unknown_not_known_absence():
    assert record("a", address="=A1")["address_present"] is None
    assert project_record(("a", "Confirmado"), {"event_id": 1, "classification_columns": {"manual": 2}}, KEY, "fictional")["address_present"] is None


def test_reconciliation_partitions_processed_records_without_deduplication():
    original = [record("a"), record("b"), record("dup"), record("dup"), record("omitted"), record(None)]
    processed = [record("a"), record("b", classification="Sospechoso", department="Other"), record("dup"), record("unexpected"), record(None)]
    before_original, before_processed = deepcopy(original), deepcopy(processed)
    result = compare(original, processed, RULE)
    counts = result["counts"]
    assert counts["processed_records"] == 5 and counts["original_records"] == 6
    assert counts["processed_unique_identifier_candidate"] == 2
    assert counts["processed_ambiguous_identifier_candidates"] == 1
    assert counts["processed_without_original_identifier_candidate"] == 1
    assert counts["processed_identifier_missing"] == 1
    assert counts["selected_original_without_processed_identifier_candidate"] == 1
    assert counts["unique_pairs_classification_changed"] == 1
    assert counts["unique_pairs_department_changed"] == 1
    assert not result["duplicates_deleted"] and not result["identifier_match_proves_same_episode"]
    assert original == before_original and processed == before_processed


def test_source_2022_like_suspected_records_remain_present_but_outside_literal_confirmed_rule():
    original = [record("a", classification="Sospechoso"), record("b", classification="Probable"), record("c", classification="Descartado")]
    result = compare(original, original, RULE)
    assert result["counts"]["literal_rule_excluded"] == 3
    assert result["counts"]["processed_unique_pair_outside_literal_rule"] == 3
    assert result["counts"]["processed_without_original_identifier_candidate"] == 0
    assert not result["clinical_confirmation_validated"]


def test_missing_date_and_address_do_not_prove_why_original_was_omitted():
    result = compare([record("absent", day=None, address=None)], [], RULE)
    assert result["counts"]["unrepresented_original_with_address_absent"] == 1
    assert result["counts"]["unrepresented_original_with_reference_date_unknown"] == 1
    assert not result["missing_dates_or_addresses_prove_manual_exclusion_reason"]


def test_aggregate_results_do_not_depend_on_secret_key_and_never_export_tokens():
    pairs = []
    for key in (KEY, b"another_fictional_ephemeral_key!"):
        original = [record("a", key=key), record("b", key=key)]
        processed = [record("a", key=key)]
        pairs.append(compare(original, processed, RULE))
    assert pairs[0] == pairs[1]
    assert "PRIVATE_EVENT" not in json.dumps(pairs)


def test_profile_only_exports_aggregated_domains_and_date_ranges():
    result = profile([record("PRIVATE_ID_A"), record("PRIVATE_ID_B", day=None, address=None)])
    assert result["records"] == 2
    assert result["classification_frequencies"]["manual"] == [{"literal_label": "Confirmado", "records": 2}]
    assert result["clinical_date_intervals"]["symptom_onset"]["valid_records"] == 1
    assert not result["observed_dates_prove_complete_coverage"]
    assert "PRIVATE_ID" not in json.dumps(result) and "PRIVATE_ADDRESS" not in json.dumps(result)


def test_dates_do_not_guess_numeric_serials_or_formula_values():
    assert date_value("2024-01-10")[1] == "valid"
    assert date_value(45000)[0] is None
    assert date_value("=TODAY()")[0] is None


def test_ambiguous_headers_are_not_arbitrarily_chosen_and_all_classification_columns_are_profiled():
    columns = resolve_columns(["ID evento", "Clasificacion manual", "Clasificacion automatica", "Provincia residencia", "Provincia residencia", "Departamento residencia"], {})
    assert "province" not in columns
    assert len(columns["classification_columns"]) == 2


def test_explicit_classification_alias_does_not_duplicate_a_physical_column():
    columns = resolve_columns(["ID evento", "Clasificacion manual"], {"classification_columns": {"manual": 2}})
    assert columns["classification_columns"] == {"manual": 2}


def test_cross_source_equal_identifiers_are_reported_without_episode_equivalence():
    from collections import Counter
    a = project_record(("PRIVATE_EVENT_ID",), {"event_id": 1}, KEY, "source_a")
    b = project_record(("PRIVATE_EVENT_ID",), {"event_id": 1}, KEY, "source_b")
    assert a["key"] != b["key"] and a["cross_source_key"] == b["cross_source_key"]
    result = cross_source_overlap({a["cross_source_key"]: Counter({"source_a": 1, "source_b": 2})})
    assert result["equal_literal_identifier_groups_across_source_mappings"] == 1
    assert result["records_in_cross_source_identifier_groups"] == 3
    assert not result["global_identifier_semantics_or_episode_equivalence_validated"]
    assert "PRIVATE_EVENT_ID" not in json.dumps(result)


@pytest.mark.parametrize("explicit", [{"event_id": 1}, {"province": 1}, {"classification_columns": {"manual": 1}}, {"document": 2}])
def test_personal_or_unsafe_column_mappings_are_rejected(explicit):
    with pytest.raises(ValueError):
        resolve_columns(["Nombre", "DNI", "Provincia residencia"], explicit)


def workbook(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Fictional"
    sheet.append(["ID evento", "Clasificacion manual", "Provincia residencia", "Departamento residencia", "Fecha inicio sintomas", "Domicilio", "Nombre", "DNI", "Latitud", "Longitud"])
    for values in rows:
        sheet.append(list(values) + ["PRIVATE_PERSON_NAME", "PRIVATE_DOCUMENT", -29.4, -66.8])
    book.save(path)
    book.close()


def spec(path):
    return {"path": str(path.resolve()), "source_version": "fictional_v1", "private_local_file_confirmed": True,
        "header_row": 1, "columns": COLUMNS, "comparison_classification_field": "manual",
        "expected_sha256": digest_path(path), "requested_period": None, "extraction_period": None,
        "extracted_on": None, "extraction_date_field": None}


@pytest.fixture
def delivery(tmp_path, monkeypatch):
    raw = tmp_path / "data/raw"
    raw.mkdir(parents=True)
    originals, processed, manual = (raw / name for name in ("original.xlsx", "processed.xlsx", "working.xlsx"))
    rows = [("PRIVATE_ID_A", "Confirmado", "La Rioja", "Capital", datetime(2024, 1, 10), "PRIVATE_ADDRESS"),
            ("PRIVATE_ID_B", "Sospechoso", "La Rioja", "Capital", None, None)]
    workbook(originals, rows)
    workbook(manual, rows)
    workbook(processed, rows[:1])
    options = {"version": VERSION, "local_aggregate_processing_authorized": True,
        "output_dir": "data/interim/official/fictional_reconciliation", "expected_processed_records": 1,
        "training_enabled": False, "generate_negatives": False,
        "sources": [{"source_id": "fictional", "manual_copy_status": "provided", "original": spec(originals),
            "manual_working": spec(manual), "processed": spec(processed), "selection_rule": RULE}]}
    monkeypatch.setattr("sigard_ml.official.selection_reconciliation_pipeline.require_branch", lambda root: None)
    monkeypatch.setattr("sigard_ml.official.selection_reconciliation_pipeline.subprocess.run", lambda *args, **kwargs: SimpleNamespace(returncode=0))
    return tmp_path, options


def test_local_excel_reading_does_not_modify_sources_or_export_patient_fields(delivery):
    root, options = delivery
    source = options["sources"][0]["original"]
    before = digest_path(root / "data/raw/original.xlsx")
    rows, metadata = read_local_role("fictional", "original", source, KEY)
    assert before == digest_path(root / "data/raw/original.xlsx")
    assert len(rows) == 2
    serialized = json.dumps(metadata)
    for secret in ("PRIVATE_PERSON_NAME", "PRIVATE_DOCUMENT", "PRIVATE_ADDRESS", "PRIVATE_ID_A", "-29.4", "-66.8"):
        assert secret not in serialized


def test_complete_pipeline_is_aggregate_only_private_reproducible_and_never_overwrites(delivery):
    root, options = delivery
    output = root / options["output_dir"]
    before = protected_snapshot(root, output)
    first = run_reconciliation(root, options)
    require_preserved(before, protected_snapshot(root, output))
    assert first["totals"] == {"original": 2, "manual_working": 2, "processed": 1}
    assert not first["row_level_outputs_written"] and not first["identifiers_exported"]
    assert len(first["sources"][0]["comparisons"]) == 3
    with pytest.raises(ValueError, match="nuevo"):
        run_reconciliation(root, options)
    second = run_reconciliation(root, {**options, "output_dir": "data/interim/official/fictional_repeat"})
    assert first == second
    for path in output.iterdir():
        if path.name != "integrity_manifest.private.json" and path.name != "configuration.private.json":
            assert digest_path(path) == digest_path(root / "data/interim/official/fictional_repeat" / path.name)
    serialized = (output / "report.private.json").read_text(encoding="utf-8")
    assert "PRIVATE_ID_A" not in serialized and "PRIVATE_PERSON_NAME" not in serialized


def test_unknown_original_path_blocks_before_opening_workbooks(delivery, monkeypatch):
    root, options = delivery
    options["sources"][0]["original"] = None
    monkeypatch.setattr("sigard_ml.official.selection_reconciliation_pipeline.read_local_role", lambda *args: pytest.fail("No abrir antes de conocer originales"))
    with pytest.raises(ValueError, match="rutas explícitas"):
        run_reconciliation(root, options)
    assert not (root / options["output_dir"]).exists()


def test_missing_identifier_column_requires_profile_instead_of_false_matches(delivery):
    root, options = delivery
    options["sources"][0]["original"]["columns"] = {"classification_columns": {"manual": 2}}
    # Header still identifies ID automatically; make ambiguity explicit with an unrecognized ID header.
    source = options["sources"][0]["original"]
    from openpyxl import load_workbook
    book = load_workbook(source["path"])
    book.active.cell(1, 1, "UnrecognizedIdentifier")
    book.save(source["path"])
    book.close()
    source["expected_sha256"] = digest_path(root / "data/raw/original.xlsx")
    with pytest.raises(ValueError, match="Esquema incompleto"):
        run_reconciliation(root, options)


@pytest.mark.parametrize("field,value", [("training_enabled", True), ("generate_negatives", True),
    ("local_aggregate_processing_authorized", False), ("output_dir", "frontend/public/data"), ("output_dir", "../outside")])
def test_unsafe_actions_or_output_locations_are_rejected(delivery, field, value):
    root, options = delivery
    with pytest.raises((ValueError, PermissionError)):
        validate_manifest({**options, field: value}, root)


def test_input_inside_workspace_must_be_git_ignored(delivery, monkeypatch):
    root, options = delivery
    monkeypatch.setattr("sigard_ml.official.selection_reconciliation_pipeline.subprocess.run", lambda *args, **kwargs: SimpleNamespace(returncode=1))
    with pytest.raises(PermissionError, match="Git"):
        validate_manifest(options, root)
